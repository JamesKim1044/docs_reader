"""Office loaders (docx / xlsx / pptx).

Default: extract text (cheap, high-signal for born-digital documents). Tables are
serialized in a readable form so the LLM keeps row/column structure. When layout
fidelity matters, ``--office-vision`` routes through LibreOffice -> PDF -> the PDF
loader (Stage A OCR / vision) instead.
"""

from __future__ import annotations

from pathlib import Path

from ..config import Settings
from ..convert import convert_to_pdf
from ..types import ContentPart, LoadedDocument
from .tables import rows_to_markdown


def _docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    lines: list[str] = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for ti, table in enumerate(d.tables, start=1):
        rows = [[c.text for c in row.cells] for row in table.rows]
        md = rows_to_markdown(rows)
        if md:
            lines.append(f"\n[table {ti}]\n{md}")
    return "\n".join(lines).strip()


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"\n[sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                chunks.append(" | ".join(cells))
    wb.close()
    return "\n".join(chunks).strip()


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    chunks: list[str] = []
    for si, slide in enumerate(prs.slides, start=1):
        chunks.append(f"\n[slide {si}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(run.text for run in para.runs)
                    if txt.strip():
                        chunks.append(txt)
            if shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                md = rows_to_markdown(rows)
                if md:
                    chunks.append(md)
    return "\n".join(chunks).strip()


_TEXT_EXTRACTORS = {
    ".docx": _docx_text,
    ".xlsx": _xlsx_text,
    ".xlsm": _xlsx_text,
    ".pptx": _pptx_text,
}


def load_legacy_office(path: str | Path, settings: Settings) -> LoadedDocument:
    """Legacy binary Office formats (.ppt/.doc/.xls) have no reliable pure-Python
    parser — convert to PDF via LibreOffice and reuse the PDF loader."""
    from .pdf import load_pdf

    path = Path(path)
    pdf_path = convert_to_pdf(path)
    doc = load_pdf(pdf_path, settings, source_label=str(path))
    doc.meta = {**doc.meta, "loader": "legacy-office->pdf", "original_ext": path.suffix.lower()}
    return doc


def load_office(path: str | Path, settings: Settings) -> LoadedDocument:
    path = Path(path)
    ext = path.suffix.lower()

    if settings.office_vision:
        # Convert to PDF and reuse the PDF loader for layout-preserving extraction.
        from .pdf import load_pdf

        pdf_path = convert_to_pdf(path)
        doc = load_pdf(pdf_path, settings, source_label=str(path))
        doc.meta = {**doc.meta, "loader": "office->pdf", "original_ext": ext}
        return doc

    extractor = _TEXT_EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(f"unsupported office type: {ext}")
    text = extractor(path)
    return LoadedDocument(
        parts=[ContentPart.from_text(text)] if text else [],
        source_path=str(path),
        route="text",
        page_count=None,
        meta={"loader": "office", "ext": ext},
    )
