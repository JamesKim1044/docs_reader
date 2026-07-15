"""Extract embedded images / figures from documents to files.

Supports PDF, docx, pptx, xlsx, HWP (via LibreOffice -> PDF), and plain image files.
Each extracted image is written to ``out_dir`` and described by :class:`ExtractedImage`
(without bytes when serialized to metadata). The bytes are kept in-memory so callers
can optionally attach figures to the vision path (``--figures-vision``).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

_EXT_BY_CONTENT_TYPE = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/jpg": "jpg",
    "image/gif": "gif",
    "image/bmp": "bmp",
    "image/tiff": "tiff",
    "image/x-emf": "emf",
    "image/webp": "webp",
}


@dataclass
class ExtractedImage:
    path: str
    page: Optional[int]
    index: int
    ext: str
    mime: str
    size_bytes: int
    data: bytes = field(repr=False, default=b"")

    def meta(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "page": self.page,
            "index": self.index,
            "ext": self.ext,
            "mime": self.mime,
            "size_bytes": self.size_bytes,
        }


def _mime_for(ext: str) -> str:
    ext = ext.lower().lstrip(".")
    return {"jpg": "image/jpeg", "jpeg": "image/jpeg"}.get(ext, f"image/{ext}")


def _save(out_dir: Path, stem: str, page: Optional[int], idx: int, data: bytes, ext: str) -> ExtractedImage:
    out_dir.mkdir(parents=True, exist_ok=True)
    tag = f"p{page}_" if page else ""
    name = f"{stem}_{tag}{idx:03d}.{ext}"
    fp = out_dir / name
    fp.write_bytes(data)
    return ExtractedImage(
        path=str(fp), page=page, index=idx, ext=ext, mime=_mime_for(ext),
        size_bytes=len(data), data=data,
    )


def _from_pdf(path: Path, out_dir: Path) -> list[ExtractedImage]:
    import fitz

    doc = fitz.open(str(path))
    stem = path.stem
    out: list[ExtractedImage] = []
    seen: set[int] = set()
    idx = 0
    try:
        for pno, page in enumerate(doc, start=1):
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                base = doc.extract_image(xref)
                data, ext = base.get("image"), base.get("ext", "png")
                if not data:
                    continue
                out.append(_save(out_dir, stem, pno, idx, data, ext))
                idx += 1
    finally:
        doc.close()
    return out


def _from_docx(path: Path, out_dir: Path) -> list[ExtractedImage]:
    import docx

    d = docx.Document(str(path))
    out: list[ExtractedImage] = []
    idx = 0
    for rel in d.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        blob = part.blob
        ext = _EXT_BY_CONTENT_TYPE.get(getattr(part, "content_type", ""), None)
        if ext is None:
            ext = Path(getattr(part, "partname", "img.png")).suffix.lstrip(".") or "png"
        out.append(_save(out_dir, path.stem, None, idx, blob, ext))
        idx += 1
    return out


def _from_pptx(path: Path, out_dir: Path) -> list[ExtractedImage]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[ExtractedImage] = []
    idx = 0
    for sno, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            try:
                image = shape.image
            except Exception:
                continue
            out.append(_save(out_dir, path.stem, sno, idx, image.blob, image.ext or "png"))
            idx += 1
    return out


def _from_xlsx(path: Path, out_dir: Path) -> list[ExtractedImage]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    out: list[ExtractedImage] = []
    idx = 0
    for ws in wb.worksheets:
        for image in getattr(ws, "_images", []) or []:
            try:
                data = image._data()  # openpyxl private accessor
            except Exception:
                continue
            ext = (getattr(image, "format", None) or "png").lower()
            out.append(_save(out_dir, path.stem, None, idx, data, ext))
            idx += 1
    wb.close()
    return out


def _from_image_file(path: Path, out_dir: Path) -> list[ExtractedImage]:
    data = path.read_bytes()
    ext = path.suffix.lstrip(".").lower() or "png"
    return [_save(out_dir, path.stem, 1, 0, data, ext)]


_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}


def extract_images(path: str | Path, out_dir: str | Path) -> list[ExtractedImage]:
    """Extract embedded images from a document, saving them under ``out_dir``."""
    path = Path(path)
    out_dir = Path(out_dir)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _from_pdf(path, out_dir)
    if ext == ".docx":
        return _from_docx(path, out_dir)
    if ext == ".pptx":
        return _from_pptx(path, out_dir)
    if ext in (".xlsx", ".xlsm"):
        return _from_xlsx(path, out_dir)
    if ext in (".hwp", ".hwpx"):
        from .convert import convert_to_pdf

        pdf = convert_to_pdf(path)
        return _from_pdf(Path(pdf), out_dir)
    if ext in _IMAGE_EXTS:
        return _from_image_file(path, out_dir)
    return []
