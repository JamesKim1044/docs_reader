"""Test fixtures — generate small sample documents on the fly."""

from __future__ import annotations

import io
from pathlib import Path

import pytest


@pytest.fixture
def text_pdf(tmp_path: Path) -> Path:
    """A born-digital PDF with a real text layer (routes to the text path)."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    lines = [
        "TAX INVOICE",
        "Invoice No: INV-2024-0091",
        "Date: 2024-03-05  Due: 2024-04-04",
        "From: Acme Supplies Ltd",
        "Bill To: Globex Corp",
        "Total: 99.00 USD",
    ]
    y = 72
    for ln in lines:
        page.insert_text((72, y), ln, fontsize=12)
        y += 18
    out = tmp_path / "invoice.pdf"
    doc.save(str(out))
    doc.close()
    return out


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    """An image-only PDF (no text layer) -> routes to the vision path."""
    import fitz
    from PIL import Image

    img = Image.new("RGB", (600, 300), "white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    doc = fitz.open()
    page = doc.new_page(width=600, height=300)
    page.insert_image(fitz.Rect(0, 0, 600, 300), stream=buf.getvalue())
    out = tmp_path / "scan.pdf"
    doc.save(str(out))
    doc.close()
    return out


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    import docx

    d = docx.Document()
    d.add_paragraph("Invoice No: INV-DOCX-1")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Total"
    table.rows[0].cells[1].text = "42.00"
    out = tmp_path / "invoice.docx"
    d.save(str(out))
    return out


@pytest.fixture
def table_image_pdf(tmp_path: Path) -> Path:
    """A born-digital PDF with a ruled table and one embedded image."""
    import io

    import fitz
    from PIL import Image

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 30), "INVOICE INV-TBL-1", fontsize=12)
    cols = [50, 150, 250, 350]
    rows = [50, 90, 130, 170]
    for x in cols:
        page.draw_line((x, rows[0]), (x, rows[-1]))
    for y in rows:
        page.draw_line((cols[0], y), (cols[-1], y))
    cells = [["Item", "Qty", "Amount"], ["Widget A", "10", "50.00"], ["Widget B", "2", "40.00"]]
    for r, ry in enumerate(rows[:-1]):
        for c, cx in enumerate(cols[:-1]):
            page.insert_text((cx + 5, ry + 25), cells[r][c], fontsize=10)
    img = Image.new("RGB", (80, 40), "blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    page.insert_image(fitz.Rect(50, 190, 130, 230), stream=buf.getvalue())
    out = tmp_path / "table_image.pdf"
    doc.save(str(out))
    doc.close()
    return out


@pytest.fixture
def docx_with_image(tmp_path: Path) -> Path:
    import io

    import docx
    from PIL import Image

    d = docx.Document()
    d.add_paragraph("Doc with figure")
    img = Image.new("RGB", (60, 30), "red")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    d.add_picture(buf)
    out = tmp_path / "with_image.docx"
    d.save(str(out))
    return out


@pytest.fixture
def schema_py() -> Path:
    return Path(__file__).resolve().parents[1] / "schemas" / "invoice.py"


@pytest.fixture
def schema_json() -> Path:
    return Path(__file__).resolve().parents[1] / "schemas" / "invoice.json"
